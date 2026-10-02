"""Étape 4 : bilan des décisions des relecteurs -> indicateurs du PoC.

Usage :
    python bilan.py "Projet_pipeline/Projet_validation.xlsx"

Lit la colonne Décision du classeur de validation et calcule :
    - l'avancement de la relecture ;
    - la précision de l'agent (part des signalements qui sont de vrais problèmes) ;
    - le taux de reprise des questions (telles quelles ou modifiées) ;
    - les oublis déclarés par les relecteurs sur les lignes OK ;
    - le détail par règle et par sévérité.
Écrit <nom>_bilan.xlsx à côté du classeur (le classeur de validation n'est pas modifié).
"""
import argparse
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from commun import DECISIONS, REGLES, cle, en_texte, ids_regles, lire_feuille

VRAIS = {"Acceptée", "Modifiée", "Rejetée"}


def pct(a, b):
    return f"{100 * a / b:.0f} %" if b else "-"


def main():
    ap = argparse.ArgumentParser(description="Calcule les indicateurs du PoC à partir des décisions des relecteurs.")
    ap.add_argument("classeur", help="fichier <nom>_validation.xlsx relu par les ingénieurs")
    a = ap.parse_args()
    chemin = Path(a.classeur)
    if not chemin.exists():
        sys.exit(f"ERREUR : fichier introuvable : {chemin}")

    wb = load_workbook(chemin, read_only=True, data_only=True)
    if "Validation" not in wb.sheetnames:
        sys.exit("ERREUR : feuille 'Validation' absente (fichier produit par controle.py ?)")
    _, lignes = lire_feuille(wb["Validation"])
    wb.close()

    dec_norm = {cle(d): d for d in DECISIONS}
    signalees, ok, autres_valeurs = [], [], Counter()
    for _, d in lignes:
        verdict = en_texte(d.get("Verdict"))
        brut = en_texte(d.get("Décision"))
        dec = dec_norm.get(cle(brut), "")
        if brut and not dec:
            autres_valeurs[brut] += 1
        x = {"ID": en_texte(d.get("ID")), "Sévérité": en_texte(d.get("Sévérité")),
             "Règles": ids_regles(en_texte(d.get("Règles"))), "Décision": dec, "Validateur": en_texte(d.get("Validateur"))}
        if verdict == "A_CLARIFIER":
            signalees.append(x)
        elif verdict == "OK":
            ok.append(x)

    decidees = [x for x in signalees if x["Décision"] and x["Décision"] != "Oubli agent"]
    c = Counter(x["Décision"] for x in decidees)
    vrais = sum(c[d] for d in VRAIS)
    oublis = sum(1 for x in ok if x["Décision"] == "Oubli agent")
    ok_relues = sum(1 for x in ok if x["Décision"])

    indicateurs = [
        ("Exigences signalées (A_CLARIFIER)", len(signalees)),
        ("  dont décidées par un relecteur", f"{len(decidees)} ({pct(len(decidees), len(signalees))})"),
        ("Précision de l'agent : vrais problèmes / décidées", pct(vrais, len(decidees))),
        ("Questions reprises telles quelles (Acceptée)", pct(c["Acceptée"], len(decidees))),
        ("Questions reprises après modification (Modifiée)", pct(c["Modifiée"], len(decidees))),
        ("Problèmes réels sans question (Rejetée)", pct(c["Rejetée"], len(decidees))),
        ("Faux positifs", pct(c["Faux positif"], len(decidees))),
        ("Lignes OK marquées par un relecteur", ok_relues),
        ("  dont oublis de l'agent (Oubli agent)", oublis),
        ("Rappel estimé : vrais problèmes / (vrais + oublis)", pct(vrais, vrais + oublis)),
    ]

    par_regle = defaultdict(Counter)
    for x in decidees:
        for rg in x["Règles"]:
            par_regle[rg][x["Décision"]] += 1
    par_sev = defaultdict(Counter)
    for x in decidees:
        par_sev[x["Sévérité"]][x["Décision"]] += 1

    # ---- console
    print(f"Bilan de {chemin.name}")
    for k, v in indicateurs:
        print(f"  {k:<55} {v}")
    if autres_valeurs:
        print("ATTENTION : valeurs de Décision non reconnues : " + ", ".join(f"'{k}' x{n}" for k, n in autres_valeurs.items()))
    print("  Le rappel estimé n'est fiable que si les relecteurs ont aussi parcouru les lignes OK.")

    # ---- fichier
    out = Workbook()
    ws = out.active
    ws.title = "Bilan"
    ws.append(["Indicateur", "Valeur"])
    for k, v in indicateurs:
        ws.append([k, v])
    ws.append(["", ""])
    ws.append(["Règle", "Décidées"] + DECISIONS[:4] + ["Taux de faux positifs"])
    for rg in REGLES:
        if rg in par_regle:
            cc = par_regle[rg]
            n = sum(cc.values())
            ws.append([f"{rg} {REGLES[rg][0]}", n] + [cc[d] for d in DECISIONS[:4]] + [pct(cc["Faux positif"], n)])
    ws.append(["", ""])
    ws.append(["Sévérité", "Décidées"] + DECISIONS[:4] + ["Taux de faux positifs"])
    for sv in ["BLOQUANT", "MAJEUR"]:
        cc = par_sev.get(sv, Counter())
        n = sum(cc.values())
        ws.append([sv, n] + [cc[d] for d in DECISIONS[:4]] + [pct(cc["Faux positif"], n)])
    for row in ws.iter_rows():
        if row[0].value in ("Indicateur", "Règle", "Sévérité"):
            for cell in row:
                cell.font = Font(bold=True)
    for col, w in zip("ABCDEFG", (55, 12, 11, 11, 11, 13, 20)):
        ws.column_dimensions[col].width = w
    sortie = chemin.with_name(chemin.name.replace("_validation", "_bilan"))
    if sortie == chemin:
        sortie = chemin.with_name(chemin.stem + "_bilan.xlsx")
    out.save(sortie)
    print(f"Bilan écrit : {sortie}")


if __name__ == "__main__":
    main()
