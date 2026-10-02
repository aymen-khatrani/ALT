"""Étape 1 : export .CUS.xlsx -> lots à déposer dans l'agent Copilot.

Usage :
    python prepare.py "Projet.CUS.xlsx"
    python prepare.py "Projet.CUS.xlsx" --taille-lot 20 --col-texte "Object Text" --types-exigence "Requirement"

Produit, à côté du fichier source, un dossier <nom>_pipeline/ contenant :
    <nom>_prepare.xlsx   fichier maître (toutes les lignes, statut, lot, signaux du lexique)
    lots/lot_01.xlsx …   un fichier par lot, colonnes ID | Texte | Signaux_lexique
    resultats/           dossier vide où déposer les fichiers renvoyés par l'agent
"""
import argparse
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from commun import (COLONNES_LOT, VERSION, charger_lexique, cle, detecter_signaux, empreinte,
                    en_texte, formater_signaux, lire_feuille)

ICI = Path(__file__).resolve().parent

NOMS_ID = ["Object Identifier", "Identifier", "ID", "Req ID", "Requirement ID", "Reference", "Ref",
           "Absolute Number", "Object Number", "Identifiant", "Référence", "Num"]
NOMS_TEXTE = ["Object Text", "Requirement Text", "Requirement", "Text", "Texte", "Exigence",
              "Statement", "Description"]
NOMS_TYPE = ["Object Type", "Requirement Type", "Type", "Category", "Catégorie", "Nature",
             "Classification", "Kind"]
MOTS_HORS = ["heading", "title", "titre", "information", "info", "comment", "note", "definition",
             "chapter", "chapitre", "section", "rationale", "figure", "table", "glossary"]
MOTS_EXIGENCE = ["requirement", "req", "exigence", "shall", "constraint", "contrainte",
                 "functional", "performance"]


def choisir_colonne(entetes, lignes, noms, forcee, role):
    if forcee:
        trouvee = next((e for e in entetes if cle(e) == cle(forcee)), None)
        if not trouvee:
            sys.exit(f"ERREUR : colonne {role} '{forcee}' introuvable. Colonnes disponibles : {entetes}")
        return trouvee
    cles = {cle(e): e for e in entetes}
    for n in noms:                       # correspondance exacte d'abord
        if cle(n) in cles:
            return cles[cle(n)]
    for n in noms:                       # puis "contient"
        for k, e in cles.items():
            if cle(n) and cle(n) in k:
                return e
    if role == "texte":                  # repli : colonne au texte moyen le plus long
        def longueur(e):
            vals = [len(en_texte(d.get(e))) for _, d in lignes]
            return sum(vals) / max(len(vals), 1)
        return max(entetes, key=longueur)
    return None


def classer_type(valeur, types_exigence):
    v = cle(valeur)
    if not v:
        return "INCONNU"
    if types_exigence is not None:
        return "EXIGENCE" if v in types_exigence else "HORS"
    if any(m in v for m in MOTS_HORS):
        return "HORS"
    if any(m in v for m in MOTS_EXIGENCE):
        return "EXIGENCE"
    return "INCONNU"


def ecrire_lot(chemin, lignes, numero):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Lot_{numero:02d}"
    ws.append(COLONNES_LOT)
    for x in lignes:
        ws.append([x["ID"], x["Texte"], x["Signaux_lexique"]])
    for c, largeur in zip("ABC", (18, 90, 60)):
        ws.column_dimensions[c].width = largeur
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for ligne in ws.iter_rows(min_row=2):
        for cellule in ligne:
            cellule.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(chemin)


def main():
    ap = argparse.ArgumentParser(description="Prépare un export .CUS.xlsx en lots pour l'agent Copilot.")
    ap.add_argument("source", help="fichier .CUS.xlsx exporté")
    ap.add_argument("--taille-lot", type=int, default=20, help="exigences par lot (défaut 20)")
    ap.add_argument("--feuille", help="nom de la feuille (défaut : celle qui contient le plus de lignes)")
    ap.add_argument("--col-id", help="nom exact de la colonne ID")
    ap.add_argument("--col-texte", help="nom exact de la colonne texte de l'exigence")
    ap.add_argument("--col-type", help="nom exact de la colonne type d'objet ('aucune' pour l'ignorer)")
    ap.add_argument("--types-exigence", help="valeurs de type à analyser, séparées par des virgules")
    ap.add_argument("--lexique", default=str(ICI / "lexique.csv"))
    ap.add_argument("--sortie", help="dossier de sortie (défaut : <nom>_pipeline à côté du fichier)")
    a = ap.parse_args()

    source = Path(a.source)
    if not source.exists():
        sys.exit(f"ERREUR : fichier introuvable : {source}")
    nom = source.name.split(".")[0]
    sortie = Path(a.sortie) if a.sortie else source.parent / f"{nom}_pipeline"
    (sortie / "lots").mkdir(parents=True, exist_ok=True)
    (sortie / "resultats").mkdir(exist_ok=True)
    for ancien in (sortie / "lots").glob("lot_*.xlsx"):
        ancien.unlink()

    lexique = charger_lexique(a.lexique)

    # --- lecture
    wb = load_workbook(source, read_only=True, data_only=True)
    if a.feuille:
        if a.feuille not in wb.sheetnames:
            sys.exit(f"ERREUR : feuille '{a.feuille}' absente. Feuilles : {wb.sheetnames}")
        ws = wb[a.feuille]
    else:
        ws = max(wb.worksheets, key=lambda w: w.max_row or 0)
    entetes, lignes = lire_feuille(ws)
    nom_feuille = ws.title
    wb.close()
    if not lignes:
        sys.exit("ERREUR : aucune ligne de données trouvée.")

    col_texte = choisir_colonne(entetes, lignes, NOMS_TEXTE, a.col_texte, "texte")
    col_id = choisir_colonne(entetes, lignes, NOMS_ID, a.col_id, "ID")
    col_type = None if (a.col_type or "").lower() == "aucune" else \
        choisir_colonne(entetes, lignes, NOMS_TYPE, a.col_type, "type")
    if col_id == col_texte:
        col_id = None
    types_exigence = {cle(t) for t in a.types_exigence.split(",")} if a.types_exigence else None

    print(f"Feuille      : {nom_feuille}")
    print(f"Colonne ID   : {col_id or '(aucune -> numéro de ligne Excel R<n>)'}")
    print(f"Colonne texte: {col_texte}")
    print(f"Colonne type : {col_type or '(aucune -> toutes les lignes avec texte sont analysées)'}")

    # --- classement et signaux
    vus, inconnus, enregistrements = Counter(), Counter(), []
    for num_ligne, d in lignes:
        texte = en_texte(d.get(col_texte))
        ident = en_texte(d.get(col_id)) if col_id else ""
        ident = ident or f"R{num_ligne}"
        vus[ident] += 1
        if vus[ident] > 1:
            ident = f"{ident}_{vus[ident]}"
        type_src = en_texte(d.get(col_type)) if col_type else ""
        if not texte:
            statut = "HORS_PERIMETRE"
        elif col_type:
            classe = classer_type(type_src, types_exigence)
            if classe == "INCONNU":
                inconnus[type_src or "(vide)"] += 1
            statut = "HORS_PERIMETRE" if classe == "HORS" else "A_ANALYSER"
        else:
            statut = "A_ANALYSER"
        signaux = formater_signaux(detecter_signaux(texte, lexique)) if statut == "A_ANALYSER" else ""
        enregistrements.append({"ID": ident, "Ligne_source": num_ligne, "Lot": "", "Statut_prepare": statut,
                                "Type_source": type_src, "Texte": texte, "Signaux_lexique": signaux})

    doublons = [k for k, n in vus.items() if n > 1]
    if doublons:
        print(f"ATTENTION : {len(doublons)} ID en double, suffixés _2, _3… (ex. {doublons[:3]})")
    if inconnus:
        print("ATTENTION : valeurs de type non reconnues, lignes analysées par précaution :")
        for v, n in inconnus.most_common():
            print(f"   - '{v}' ({n} lignes)")
        print("   -> pour les exclure, relancez avec --types-exigence \"<valeurs à analyser>\"")

    # --- lots
    a_analyser = [x for x in enregistrements if x["Statut_prepare"] == "A_ANALYSER"]
    if not a_analyser:
        sys.exit("ERREUR : aucune exigence à analyser (vérifiez --col-texte / --types-exigence).")
    taille = max(1, a.taille_lot)
    nb_lots = (len(a_analyser) + taille - 1) // taille
    for i in range(nb_lots):
        paquet = a_analyser[i * taille:(i + 1) * taille]
        for x in paquet:
            x["Lot"] = i + 1
        ecrire_lot(sortie / "lots" / f"lot_{i + 1:02d}.xlsx", paquet, i + 1)

    # --- fichier maître
    nb_signaux = sum(x["Signaux_lexique"].count(":") for x in a_analyser)
    signature = (f"SIGNATURE prepare v{VERSION} | source={source.name} | {len(enregistrements)} lignes | "
                 f"{len(a_analyser)} a analyser | {len(enregistrements) - len(a_analyser)} hors perimetre | "
                 f"{nb_lots} lots de {taille} max | {nb_signaux} regles signalees | lexique={empreinte(a.lexique)}")
    m = Workbook()
    ws = m.active
    ws.title = "Prepare"
    cols = ["ID", "Ligne_source", "Lot", "Statut_prepare", "Type_source", "Texte", "Signaux_lexique"]
    ws.append(cols)
    for x in enregistrements:
        ws.append([x[c] for c in cols])
    for c, largeur in zip("ABCDEFG", (18, 12, 6, 16, 16, 80, 60)):
        ws.column_dimensions[c].width = largeur
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F3864")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    info = m.create_sheet("Info")
    for k, v in [("signature", signature), ("source", str(source.resolve())), ("date", datetime.now().isoformat(timespec="seconds")),
                 ("feuille", nom_feuille), ("col_id", col_id or ""), ("col_texte", col_texte), ("col_type", col_type or ""),
                 ("taille_lot", taille), ("nb_lots", nb_lots), ("lexique", str(Path(a.lexique).resolve())),
                 ("lexique_empreinte", empreinte(a.lexique)), ("version", VERSION)]:
        info.append([k, v])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 120
    maitre = sortie / f"{nom}_prepare.xlsx"
    m.save(maitre)

    print()
    print(signature)
    print(f"Fichier maître : {maitre}")
    print(f"Lots           : {sortie / 'lots'}  ({nb_lots} fichiers)")
    print(f"Résultats agent: déposez les fichiers resultats_lot_XX.xlsx dans {sortie / 'resultats'}")


if __name__ == "__main__":
    main()
