"""Étape 3 : résultats de l'agent -> contrôles -> classeur de validation pour les ingénieurs.

Usage :
    python controle.py "Projet_pipeline"
    python controle.py "Projet_pipeline" --reference reference.xlsx

Lit le fichier maître produit par prepare.py et tous les fichiers renvoyés par l'agent
(resultats/*.xlsx, ou *.json / *.txt contenant le tableau JSON), puis vérifie que l'agent n'a
rien oublié ni inventé. Produit <nom>_validation.xlsx. Si ce fichier existe déjà, les décisions
humaines déjà saisies sont reprises (ancienne version sauvegardée en .bak.xlsx).
"""
import argparse
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from commun import (COLONNES_RESULTAT, CONFIANCES, DECISIONS, REGLES, VERSION, cle, en_texte, ids_regles,
                    lire_feuille, nombres, nombres_hors_placeholders, texte_normalise)

# Noms acceptés pour chaque colonne de résultat (l'agent peut varier l'orthographe)
ALIAS = {
    "ID": ["id", "identifier"],
    "Verdict": ["verdict"],
    "Sévérité": ["severite", "severity", "sev"],
    "Règles": ["regles", "rules", "regle", "rule"],
    "Signaux écartés": ["signauxecartes", "ecartes", "signauxrejetes", "dismissedsignals", "dismissed"],
    "Extrait": ["extrait", "extraits", "excerpt", "excerpts"],
    "Explication": ["explication", "explanation"],
    "Question client": ["questionclient", "question", "clarificationquestion", "customerquestion"],
    "Formulation proposée": ["formulationproposee", "formulation", "proposedwording", "wording"],
    "Confiance": ["confiance", "confidence"],
}
CLE_VERS_COL = {a: col for col, alias in ALIAS.items() for a in alias}
RE_CANDIDAT_REGLE = re.compile(r"\b[A-Z]{3}-\d{2}\b")
RE_FRANCAIS = re.compile(r"\b(le|la|les|des|du|une|est|sont|doit|pour|avec|dans|qui|que|cette|aux)\b", re.I)
HUMAIN = ["Décision", "Question finale", "Commentaire", "Validateur"]

# Couleurs
GRIS, BLEU, ORANGE, VERT = "D9D9D9", "DDEBF7", "FCE4D6", "E2EFDA"
FOND_SEV = {"BLOQUANT": "F8CBAD", "MAJEUR": "FFE699", "MINEUR": "FFF2CC"}


# ---------------------------------------------------------------- lecture
def lire_maitre(dossier):
    candidats = sorted(dossier.glob("*_prepare.xlsx"))
    if len(candidats) != 1:
        sys.exit(f"ERREUR : il faut exactement un fichier *_prepare.xlsx dans {dossier} (trouvés : {len(candidats)})")
    wb = load_workbook(candidats[0], read_only=True, data_only=True)
    _, lignes = lire_feuille(wb["Prepare"])
    info = {en_texte(r[0]): en_texte(r[1]) for r in wb["Info"].iter_rows(values_only=True)} if "Info" in wb.sheetnames else {}
    wb.close()
    maitre = [{k: en_texte(v) for k, v in d.items()} for _, d in lignes]
    return candidats[0], maitre, info


def canoniser(d):
    sortie = {}
    for k, v in d.items():
        col = CLE_VERS_COL.get(cle(k))
        if col and col not in sortie:
            sortie[col] = en_texte(v) if not isinstance(v, list) else ", ".join(map(str, v))
    return sortie


def lire_fichier_resultat(chemin):
    """-> liste de dicts aux colonnes canoniques."""
    if chemin.suffix.lower() == ".xlsx":
        wb = load_workbook(chemin, read_only=True, data_only=True)
        ws = max(wb.worksheets, key=lambda w: w.max_row or 0)
        _, lignes = lire_feuille(ws)
        wb.close()
        return [canoniser(d) for _, d in lignes]
    brut = chemin.read_text(encoding="utf-8-sig", errors="replace")
    debut, fin = brut.find("["), brut.rfind("]")
    if debut < 0 or fin < debut:
        raise ValueError("aucun tableau JSON trouvé")
    donnees = json.loads(brut[debut:fin + 1])
    return [canoniser(d) for d in donnees if isinstance(d, dict)]


def lire_resultats(dossier_res):
    fichiers = sorted([f for f in dossier_res.glob("*") if f.suffix.lower() in (".xlsx", ".json", ".txt")
                       and not f.name.startswith("~$")], key=lambda f: f.stat().st_mtime)
    par_id, anomalies_fichier = {}, []
    for f in fichiers:
        try:
            lignes = lire_fichier_resultat(f)
        except Exception as e:  # fichier illisible : on le signale et on continue
            anomalies_fichier.append(("", "", "FICHIER_ILLISIBLE", f"{f.name} : {e}"))
            continue
        manquantes = [c for c in COLONNES_RESULTAT if lignes and c not in lignes[0]]
        if manquantes:
            anomalies_fichier.append(("", "", "COLONNE_MANQUANTE", f"{f.name} : {', '.join(manquantes)}"))
        vus = set()
        for d in lignes:
            ident = d.get("ID", "")
            if not ident:
                continue
            if ident in vus:
                anomalies_fichier.append((ident, "", "DOUBLON", f"{f.name} : ID présent plusieurs fois, 1re occurrence gardée"))
                continue
            vus.add(ident)
            d["_fichier"] = f.name
            par_id[ident] = d  # un fichier plus récent remplace un plus ancien (relance d'un lot)
    return fichiers, par_id, anomalies_fichier


# ---------------------------------------------------------------- contrôles
def normaliser_verdict(v):
    return {"aclarifier": "A_CLARIFIER", "ok": "OK", "horsperimetre": "HORS_PERIMETRE"}.get(cle(v), v.strip())


def normaliser_severite(v):
    return {"bloquant": "BLOQUANT", "majeur": "MAJEUR", "mineur": "MINEUR", "": ""}.get(cle(v), v.strip())


def regles_et_inconnues(texte):
    candidats = RE_CANDIDAT_REGLE.findall(texte or "")
    return ids_regles(texte), sorted({c for c in candidats if c not in REGLES})


def signaux_du_lexique(s):
    return set(ids_regles(s))


def controler(m, r):
    """m = ligne du maître, r = résultat de l'agent. -> (résultat normalisé, liste d'anomalies (type, détail))."""
    an = []
    verdict = normaliser_verdict(r.get("Verdict", ""))
    sev = normaliser_severite(r.get("Sévérité", ""))
    regles, inc1 = regles_et_inconnues(r.get("Règles", ""))
    ecartes, inc2 = regles_et_inconnues(r.get("Signaux écartés", ""))
    question, formulation = r.get("Question client", ""), r.get("Formulation proposée", "")
    texte = m["Texte"]

    if verdict not in ("A_CLARIFIER", "OK", "HORS_PERIMETRE"):
        an.append(("VERDICT_INVALIDE", f"'{r.get('Verdict', '')}'"))
    if sev not in ("BLOQUANT", "MAJEUR", "MINEUR", ""):
        an.append(("SEVERITE_INVALIDE", f"'{r.get('Sévérité', '')}'"))
    if inc1 or inc2:
        an.append(("REGLE_INCONNUE", ", ".join(inc1 + inc2)))

    if verdict == "A_CLARIFIER" and sev not in ("BLOQUANT", "MAJEUR"):
        an.append(("INCOHERENCE", f"A_CLARIFIER avec sévérité '{sev}'"))
    if verdict == "OK" and sev in ("BLOQUANT", "MAJEUR"):
        an.append(("INCOHERENCE", f"OK avec sévérité {sev}"))
    if verdict in ("A_CLARIFIER", "OK") and regles and not sev:
        an.append(("INCOHERENCE", "règles listées sans sévérité"))
    if verdict == "A_CLARIFIER" and not regles:
        an.append(("INCOHERENCE", "A_CLARIFIER sans règle"))

    if verdict == "A_CLARIFIER":
        if not question:
            an.append(("CHAMP_MANQUANT", "Question client vide"))
        if not formulation:
            an.append(("CHAMP_MANQUANT", "Formulation proposée vide"))
    elif question or formulation:
        an.append(("CHAMP_EN_TROP", f"question/formulation sur une ligne {verdict}"))

    signales = signaux_du_lexique(m.get("Signaux_lexique", ""))
    non_arb = sorted(signales - set(regles) - set(ecartes))
    if non_arb and verdict != "HORS_PERIMETRE":
        an.append(("SIGNAL_NON_ARBITRE", ", ".join(non_arb)))
    doubles = sorted(set(regles) & set(ecartes))
    if doubles:
        an.append(("SIGNAL_DOUBLE", ", ".join(doubles)))

    texte_n = texte_normalise(texte)
    for e in [x.strip().strip('"').strip("'").strip("«»“”").strip() for x in re.split(r"\s;\s|;", r.get("Extrait", ""))]:
        if e and texte_normalise(e) not in texte_n:
            an.append(("EXTRAIT_INTROUVABLE", f"'{e}'"))

    ajoutes = sorted(nombres_hors_placeholders(formulation) - nombres(texte))
    if ajoutes:
        an.append(("CHIFFRE_AJOUTE", ", ".join(ajoutes)))

    for nom, val in (("Question client", question), ("Formulation proposée", formulation)):
        if len(RE_FRANCAIS.findall(val)) >= 3:
            an.append(("LANGUE", f"{nom} semble en français"))

    conf = r.get("Confiance", "")
    conf_n = next((c for c in CONFIANCES if cle(c) == cle(conf)), conf)
    if conf and conf_n not in CONFIANCES:
        an.append(("CONFIANCE_INVALIDE", f"'{conf}'"))

    propre = {"Verdict": verdict, "Sévérité": sev, "Règles": ", ".join(regles), "Signaux écartés": ", ".join(ecartes),
              "Extrait": r.get("Extrait", ""), "Explication": r.get("Explication", ""), "Question client": question,
              "Formulation proposée": formulation, "Confiance": conf_n, "_fichier": r.get("_fichier", "")}
    return propre, an


# ---------------------------------------------------------------- évaluation
def evaluer(reference_path, lignes):
    wb = load_workbook(reference_path, read_only=True, data_only=True)
    _, ref = lire_feuille(wb.active)
    wb.close()
    attendus = {}
    for _, d in ref:
        dd = {cle(k): en_texte(v) for k, v in d.items()}
        ident = dd.get("id", "")
        if ident:
            attendus[ident] = (normaliser_verdict(dd.get("verdictattendu", dd.get("verdict", ""))),
                               set(ids_regles(dd.get("reglesattendues", dd.get("regles", "")))))
    obtenus = {x["ID"]: x for x in lignes}
    communs = [i for i in attendus if i in obtenus and obtenus[i]["Verdict"] not in ("", "NON_ANALYSE")]
    classes = ["A_CLARIFIER", "OK", "HORS_PERIMETRE"]
    matrice = {a: Counter() for a in classes}
    tp = fp = fn = tn = 0
    par_regle = defaultdict(lambda: [0, 0, 0])
    for i in communs:
        va, ra = attendus[i]
        vo = obtenus[i]["Verdict"]
        ro = set(ids_regles(obtenus[i]["Règles"]))
        if va in matrice:
            matrice[va][vo] += 1
        pos_a, pos_o = va == "A_CLARIFIER", vo == "A_CLARIFIER"
        tp += pos_a and pos_o
        fp += (not pos_a) and pos_o
        fn += pos_a and not pos_o
        tn += (not pos_a) and not pos_o
        for rg in ro & ra:
            par_regle[rg][0] += 1
        for rg in ro - ra:
            par_regle[rg][1] += 1
        for rg in ra - ro:
            par_regle[rg][2] += 1
    div = lambda a, b: round(a / b, 3) if b else None
    TP = sum(v[0] for v in par_regle.values())
    FP = sum(v[1] for v in par_regle.values())
    FN = sum(v[2] for v in par_regle.values())
    return {
        "n": len(communs), "n_reference": len(attendus),
        "exactitude_verdict": div(sum(matrice[c][c] for c in classes), len(communs)),
        "detection": {"VP": tp, "FP": fp, "FN": fn, "VN": tn, "precision": div(tp, tp + fp), "rappel": div(tp, tp + fn)},
        "regles": {"VP": TP, "FP": FP, "FN": FN, "precision": div(TP, TP + FP), "rappel": div(TP, TP + FN)},
        "par_regle": {k: par_regle[k] for k in sorted(par_regle)}, "matrice": matrice, "classes": classes,
    }


# ---------------------------------------------------------------- écriture
def entete(ws, colonnes, couleurs):
    ws.append(colonnes)
    for c, coul in zip(ws[1], couleurs):
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=coul)
        c.alignment = Alignment(wrap_text=True, vertical="center")


def largeurs(ws, valeurs):
    for i, l in enumerate(valeurs, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = l


def reprendre_decisions(chemin):
    if not chemin.exists():
        return {}
    wb = load_workbook(chemin, read_only=True, data_only=True)
    if "Validation" not in wb.sheetnames:
        return {}
    _, lignes = lire_feuille(wb["Validation"])
    wb.close()
    return {en_texte(d.get("ID")): {h: en_texte(d.get(h)) for h in HUMAIN} for _, d in lignes
            if any(en_texte(d.get(h)) for h in HUMAIN)}


def main():
    ap = argparse.ArgumentParser(description="Contrôle les résultats de l'agent et produit le classeur de validation.")
    ap.add_argument("dossier", help="dossier <nom>_pipeline créé par prepare.py")
    ap.add_argument("--resultats", help="dossier des fichiers de l'agent (défaut : <dossier>/resultats)")
    ap.add_argument("--reference", help="fichier de référence annoté (ID | Verdict_attendu | Règles_attendues)")
    a = ap.parse_args()

    dossier = Path(a.dossier)
    chemin_maitre, maitre, info = lire_maitre(dossier)
    dossier_res = Path(a.resultats) if a.resultats else dossier / "resultats"
    if not dossier_res.exists():
        sys.exit(f"ERREUR : dossier de résultats introuvable : {dossier_res}")
    fichiers, par_id, anomalies = lire_resultats(dossier_res)

    ids_maitre = {x["ID"] for x in maitre}
    for ident, r in par_id.items():
        if ident not in ids_maitre:
            anomalies.append((ident, "", "ID_INCONNU", f"{r['_fichier']} : ID absent du fichier maître"))

    lignes, lots_a_relancer = [], set()
    for m in maitre:
        x = {"ID": m["ID"], "Lot": m.get("Lot", ""), "Texte": m["Texte"], "Signaux lexique": m.get("Signaux_lexique", "")}
        if m.get("Statut_prepare") != "A_ANALYSER":
            x.update({"Verdict": "HORS_PERIMETRE", "Contrôles": "classé hors périmètre par prepare.py"})
        elif m["ID"] not in par_id:
            x.update({"Verdict": "NON_ANALYSE", "Contrôles": "MANQUANT"})
            anomalies.append((m["ID"], x["Lot"], "MANQUANT", "aucun résultat de l'agent"))
            lots_a_relancer.add(x["Lot"])
        else:
            propre, an = controler(m, par_id[m["ID"]])
            x.update(propre)
            x["Contrôles"] = " | ".join(f"{t}: {d}" for t, d in an)
            for t, d in an:
                anomalies.append((m["ID"], x["Lot"], t, d))
                if t == "VERDICT_INVALIDE":
                    lots_a_relancer.add(x["Lot"])
        lignes.append(x)

    # ---- classeur
    nom = chemin_maitre.name.replace("_prepare.xlsx", "")
    sortie = dossier / f"{nom}_validation.xlsx"
    decisions = reprendre_decisions(sortie)
    if sortie.exists():
        sauvegarde = dossier / f"{nom}_validation_{datetime.now():%Y%m%d_%H%M%S}.bak.xlsx"
        shutil.copy2(sortie, sauvegarde)

    wb = Workbook()
    ws = wb.active
    ws.title = "Validation"
    cols = (["ID", "Lot", "Texte", "Signaux lexique"] + ["Verdict", "Sévérité", "Règles", "Signaux écartés", "Extrait",
            "Explication", "Question client", "Formulation proposée", "Confiance"] + ["Contrôles"] + HUMAIN)
    couleurs = [GRIS] * 4 + [BLEU] * 9 + [ORANGE] + [VERT] * 4
    entete(ws, cols, couleurs)
    for x in lignes:
        x.update(decisions.get(x["ID"], {}))
        ws.append([x.get(c, "") for c in cols])
    largeurs(ws, [14, 5, 60, 30, 14, 11, 18, 14, 25, 40, 50, 60, 10, 35, 13, 50, 30, 14])
    i_sev, i_ctrl, i_dec = cols.index("Sévérité") + 1, cols.index("Contrôles") + 1, cols.index("Décision") + 1
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        sev = row[i_sev - 1].value
        if sev in FOND_SEV:
            row[i_sev - 1].fill = PatternFill("solid", fgColor=FOND_SEV[sev])
        if row[i_ctrl - 1].value and row[4].value != "HORS_PERIMETRE":
            row[i_ctrl - 1].font = Font(color="C00000")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    if lignes:
        lettre = ws.cell(row=1, column=i_dec).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{lettre}2:{lettre}{len(lignes) + 1}")

    # ---- synthèse
    analysees = [x for x in lignes if x["Verdict"] in ("A_CLARIFIER", "OK", "HORS_PERIMETRE") and x.get("_fichier")]
    verdicts = Counter(x["Verdict"] for x in lignes)
    sevs = Counter(x.get("Sévérité", "") for x in analysees if x.get("Sévérité"))
    signature = (f"SIGNATURE controle v{VERSION} | {len(fichiers)} fichiers agent | {len(analysees)} analysees | "
                 f"{verdicts.get('NON_ANALYSE', 0)} manquantes | {len(anomalies)} anomalies | "
                 f"lots a relancer: {', '.join(map(str, sorted(lots_a_relancer, key=lambda v: int(v or 0)))) or 'aucun'}")
    s = wb.create_sheet("Synthèse")
    s.append(["Élément", "Valeur"])
    for c in s[1]:
        c.font = Font(bold=True)
    for k, v in [("Signature prepare", info.get("signature", "")), ("Signature controle", signature),
                 ("Date du contrôle", datetime.now().isoformat(timespec="seconds")), ("", "")]:
        s.append([k, v])
    s.append(["Verdicts", ""])
    for v in ["A_CLARIFIER", "OK", "HORS_PERIMETRE", "NON_ANALYSE"]:
        s.append([f"  {v}", verdicts.get(v, 0)])
    s.append(["Sévérités (lignes analysées)", ""])
    for v in ["BLOQUANT", "MAJEUR", "MINEUR"]:
        s.append([f"  {v}", sevs.get(v, 0)])
    s.append(["", ""])
    s.append(["Règle", "Confirmées par l'agent", "dont signalées par le lexique", "trouvées par l'agent seul",
              "Signaux lexique", "Signaux écartés", "Taux de confirmation du lexique"])
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    for rg, (nom_rg, _) in REGLES.items():
        conf = [x for x in analysees if rg in ids_regles(x.get("Règles", ""))]
        sig = [x for x in analysees if rg in ids_regles(x.get("Signaux lexique", ""))]
        conf_sig = [x for x in conf if rg in ids_regles(x.get("Signaux lexique", ""))]
        ecart = [x for x in sig if rg in ids_regles(x.get("Signaux écartés", ""))]
        taux = round(len(conf_sig) / len(sig), 2) if sig else ""
        s.append([f"{rg} {nom_rg}", len(conf), len(conf_sig), len(conf) - len(conf_sig), len(sig), len(ecart), taux])
    s.append(["", ""])
    s.append(["Anomalies par type", ""])
    for t, n in Counter(t for _, _, t, _ in anomalies).most_common():
        s.append([f"  {t}", n])
    s.append(["Lots à relancer", ", ".join(map(str, sorted(lots_a_relancer, key=lambda v: int(v or 0)))) or "aucun"])
    largeurs(s, [42, 24, 26, 24, 16, 16, 30])

    an_ws = wb.create_sheet("Anomalies")
    entete(an_ws, ["ID", "Lot", "Type", "Détail"], [ORANGE] * 4)
    for ident, lot, t, d in anomalies:
        an_ws.append([ident, lot, t, d])
    largeurs(an_ws, [16, 6, 24, 90])
    an_ws.auto_filter.ref = an_ws.dimensions

    ev = None
    if a.reference:
        ev = evaluer(a.reference, lignes)
        e = wb.create_sheet("Évaluation")
        e.append(["Mesure", "Valeur"])
        for c in e[1]:
            c.font = Font(bold=True)
        e.append(["Exigences comparées", f"{ev['n']} / {ev['n_reference']} de la référence"])
        e.append(["Exactitude du verdict", ev["exactitude_verdict"]])
        d = ev["detection"]
        e.append(["Détection A_CLARIFIER : précision", d["precision"]])
        e.append(["Détection A_CLARIFIER : rappel", d["rappel"]])
        e.append(["Détection : VP / FP / FN / VN", f"{d['VP']} / {d['FP']} / {d['FN']} / {d['VN']}"])
        rg = ev["regles"]
        e.append(["Règles : précision", rg["precision"]])
        e.append(["Règles : rappel", rg["rappel"]])
        e.append(["", ""])
        e.append(["Matrice (lignes = attendu, colonnes = obtenu)"] + ev["classes"])
        for cl in ev["classes"]:
            e.append([cl] + [ev["matrice"][cl][o] for o in ev["classes"]])
        e.append(["", ""])
        e.append(["Règle", "VP", "FP", "FN"])
        for k, (vp, fp, fn) in ev["par_regle"].items():
            e.append([k, vp, fp, fn])
        largeurs(e, [46, 16, 16, 18])

    try:
        wb.save(sortie)
    except PermissionError:
        sortie = dossier / f"{nom}_validation_{datetime.now():%H%M%S}.xlsx"
        wb.save(sortie)
        print("ATTENTION : le classeur était ouvert dans Excel, résultat écrit sous un autre nom.")

    # ---- console
    print(signature)
    print(f"Verdicts : " + ", ".join(f"{k}={verdicts.get(k, 0)}" for k in ["A_CLARIFIER", "OK", "HORS_PERIMETRE", "NON_ANALYSE"]))
    if anomalies:
        print("Anomalies : " + ", ".join(f"{t}={n}" for t, n in Counter(t for _, _, t, _ in anomalies).most_common()))
    if decisions:
        print(f"Décisions humaines reprises : {len(decisions)}")
    if ev:
        d, rg = ev["detection"], ev["regles"]
        print(f"Évaluation ({ev['n']} exigences) : détection précision={d['precision']} rappel={d['rappel']} | "
              f"règles précision={rg['precision']} rappel={rg['rappel']}")
    print(f"Classeur de validation : {sortie}")


if __name__ == "__main__":
    main()
